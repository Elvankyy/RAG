from __future__ import annotations

import io

import httpx
from docx import Document

from tests.conftest import login


def auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


async def token_for(client: httpx.AsyncClient, username: str, password: str) -> str:
    response = await login(client, username, password)
    assert response.status_code == 200, response.text
    return response.json()["access_token"]


def make_docx(text: str) -> bytes:
    document = Document()
    document.add_heading("测试 Word 文档", level=1)
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_unit_crud_version_and_permissions(client: httpx.AsyncClient):
    token = await token_for(client, "km", "Km@12345")
    headers = auth(token)
    content = "# 测试知识单元\n\n这是第一段内容，用于验证切片与版本。\n\n## 第二节\n\n这是第二段内容。"

    created = await client.post(
        "/api/knowledge/units",
        headers=headers,
        json={"title": "测试知识单元", "content": content, "category": "测试分类", "tags": "测试"},
    )
    assert created.status_code == 201
    unit = created.json()
    unit_id = unit["id"]
    assert unit["unit_code"].startswith("KU-")
    assert unit["version"] == 1
    assert unit["chunk_count"] >= 1
    assert unit["permission_summary"] == "未配置（无访问权限）"

    detail = await client.get(f"/api/knowledge/units/{unit_id}", headers=headers)
    assert detail.status_code == 200
    assert detail.json()["content"] == content

    updated = await client.put(
        f"/api/knowledge/units/{unit_id}",
        headers=headers,
        json={"content": content + "\n\n新增第三段。", "title": "测试知识单元V2"},
    )
    assert updated.status_code == 200
    assert updated.json()["version"] == 2

    versions = await client.get(f"/api/knowledge/units/{unit_id}/versions", headers=headers)
    assert versions.status_code == 200
    assert len(versions.json()["items"]) == 2
    assert versions.json()["items"][0]["is_current"] is True

    rolled_back = await client.post(
        f"/api/knowledge/units/{unit_id}/versions/1/rollback", headers=headers
    )
    assert rolled_back.status_code == 200
    assert rolled_back.json()["version"] == 1
    assert rolled_back.json()["content"] == content

    permissions = await client.post(
        f"/api/knowledge/units/{unit_id}/permissions",
        headers=headers,
        json={
            "permissions": [
                {"target_type": "global", "target_id": None},
                {"target_type": "department", "target_id": 2},
                {"target_type": "role", "target_id": 3},
                {"target_type": "user", "target_id": 3},
            ]
        },
    )
    assert permissions.status_code == 200
    names = {item["target_name"] for item in permissions.json()["items"]}
    assert "全局公开" in names
    assert "技术部" in names
    assert "普通用户" in names
    assert "陈思远" in names

    detail = await client.get(f"/api/knowledge/units/{unit_id}", headers=headers)
    assert "全局公开" in detail.json()["permission_summary"]

    deleted = await client.request(
        "DELETE",
        "/api/knowledge/units",
        headers=headers,
        json={"ids": [unit_id]},
    )
    assert deleted.status_code == 200
    assert deleted.json()["deleted"] == 1


async def test_unit_list_and_filters(client: httpx.AsyncClient):
    token = await token_for(client, "km", "Km@12345")
    headers = auth(token)

    for index in range(3):
        await client.post(
            "/api/knowledge/units",
            headers=headers,
            json={
                "title": f"筛选测试 {index}",
                "content": f"# 筛选测试 {index}\n\n报销与差旅内容 {index}",
                "category": "筛选分类" if index < 2 else "其他分类",
            },
        )

    listed = await client.get(
        "/api/knowledge/units", headers=headers, params={"keyword": "筛选测试", "page_size": 10}
    )
    assert listed.json()["total"] == 3

    filtered = await client.get(
        "/api/knowledge/units",
        headers=headers,
        params={"keyword": "筛选测试", "category": "筛选分类"},
    )
    assert filtered.json()["total"] == 2

    categories = await client.get("/api/knowledge/categories", headers=headers)
    assert categories.status_code == 200
    assert any(item["name"] == "筛选分类" for item in categories.json()["items"])


async def test_import_md_txt_docx(client: httpx.AsyncClient):
    token = await token_for(client, "km", "Km@12345")
    headers = auth(token)

    files = [
        ("files", ("导入测试.md", ("# 导入测试" + chr(10) + chr(10) + "Markdown 内容").encode(), "text/markdown")),
        ("files", ("导入测试.txt", "纯文本导入内容".encode(), "text/plain")),
        ("files", ("导入测试.docx", make_docx("Word 文档导入内容"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
    ]
    response = await client.post(
        "/api/knowledge/import",
        headers=headers,
        data={"category": "导入分类", "tags": "导入,测试"},
        files=files,
    )
    assert response.status_code == 201, response.text
    batch = response.json()
    assert batch["total"] == 3
    assert batch["success_count"] == 3
    assert batch["failed_count"] == 0

    fetched = await client.get(
        f"/api/knowledge/import/batches/{batch['batch_id']}", headers=headers
    )
    assert fetched.status_code == 200
    assert len(fetched.json()["tasks"]) == 3
    assert all(task["status"] == "success" for task in fetched.json()["tasks"])

    listed = await client.get(
        "/api/knowledge/units", headers=headers, params={"keyword": "导入测试", "page_size": 10}
    )
    assert listed.json()["total"] >= 3


async def test_import_rejects_unsupported_and_retry(client: httpx.AsyncClient):
    token = await token_for(client, "km", "Km@12345")
    headers = auth(token)

    response = await client.post(
        "/api/knowledge/import",
        headers=headers,
        data={"category": "导入分类"},
        files=[("files", ("bad.exe", b"binary", "application/octet-stream"))],
    )
    assert response.status_code == 201
    batch = response.json()
    assert batch["failed_count"] == 1
    task = batch["tasks"][0]
    assert task["status"] == "failed"

    retry = await client.post(
        f"/api/knowledge/import/tasks/{task['id']}/retry", headers=headers
    )
    assert retry.status_code == 200

    fetched = await client.get(
        f"/api/knowledge/import/batches/{batch['batch_id']}", headers=headers
    )
    assert fetched.json()["tasks"][0]["status"] == "success"
    assert fetched.json()["success_count"] == 1


async def test_knowledge_rbac(client: httpx.AsyncClient):
    alice = await token_for(client, "alice", "User@123")
    response = await client.get("/api/knowledge/units", headers=auth(alice))
    assert response.status_code == 403
