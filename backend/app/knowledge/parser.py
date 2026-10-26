from __future__ import annotations

import io
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".md", ".markdown", ".docx", ".txt"}


class ParseError(Exception):
    pass


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"PDF 解析失败：{exc}") from exc
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            pages.append(f"第 {index} 页\n{text.strip()}")
    if not pages:
        raise ParseError("PDF 中未提取到可用文本（可能为扫描件，需开启 OCR）")
    return "\n\n".join(pages)


def _parse_docx(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Word 文档解析失败：{exc}") from exc
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    if not blocks:
        raise ParseError("Word 文档中未提取到可用文本")
    return "\n\n".join(blocks)


def parse_document(file_name: str, data: bytes) -> tuple[str, str]:
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ParseError(f"不支持的文件格式：{suffix or '未知'}，仅支持 PDF / DOCX / MD / TXT")
    if suffix == ".pdf":
        return _parse_pdf(data), "pdf"
    if suffix == ".docx":
        return _parse_docx(data), "docx"
    text = _decode(data)
    if not text.strip():
        raise ParseError("文件内容为空")
    file_type = "md" if suffix in {".md", ".markdown"} else "txt"
    return text, file_type


def derive_title(file_name: str, content: str) -> str:
    stem = Path(file_name).stem.strip()
    for line in content.splitlines():
        candidate = line.strip().lstrip("#").strip()
        if candidate and line.strip().startswith("#") and len(candidate) <= 80:
            return candidate
    return stem or "未命名知识单元"


def derive_category(title: str, content: str) -> str:
    head = f"{title}\n{content[:400]}"
    rules = [
        ("人力资源", ["薪酬", "绩效", "招聘", "考勤", "假期", "员工手册", "入职", "培训"]),
        ("财务制度", ["报销", "财务", "发票", "预算", "付款", "差旅"]),
        ("技术规范", ["技术", "研发", "代码", "架构", "接口", "安全", "运维", "部署"]),
        ("市场运营", ["市场", "营销", "品牌", "推广", "客户", "活动", "销售"]),
        ("产品管理", ["产品", "需求", "路线图", "版本", "用户研究"]),
        ("行政管理", ["行政", "办公", "会议室", "用品"]),
    ]
    for category, keywords in rules:
        if any(keyword in head for keyword in keywords):
            return category
    return "通用知识"


def clean_title(value: str) -> str:
    value = re.sub(r"[^\w一-鿿\s（）()·\-]", "", value).strip()
    return value[:120] or "未命名知识单元"
